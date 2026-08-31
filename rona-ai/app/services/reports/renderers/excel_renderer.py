"""Excel (xlsx) report renderer."""

from __future__ import annotations

import io
from typing import TYPE_CHECKING

from app.services.reports.models import ReportDocument, ReportTable
from app.services.reports.renderers.cells import (
    excel_cell,
    excel_column_letter,
    unique_sheet_name,
)

if TYPE_CHECKING:  # pragma: no cover - typing only
    from openpyxl.styles import Font
    from openpyxl.worksheet.worksheet import Worksheet

_DETAIL_COLUMN_WIDTH = 18
_SUMMARY_FIRST_DATA_ROW = 8


def _write_header_block(worksheet: "Worksheet", document: ReportDocument, title_font: "Font") -> None:
    worksheet["A1"] = document.title
    worksheet["A1"].font = title_font
    worksheet["A2"] = f"Tenant: {document.tenant_label}"
    worksheet["A3"] = f"Period: {document.period_label}"
    worksheet["A4"] = f"Generated: {document.generated_at.isoformat()}"
    worksheet["A5"] = f"Source: {document.source_system}"


def _write_summary_block(
    worksheet: "Worksheet", document: ReportDocument, bold: "Font"
) -> None:
    worksheet["A7"] = "Metric"
    worksheet["B7"] = "Value"
    worksheet["A7"].font = bold
    worksheet["B7"].font = bold
    for offset, (label, value) in enumerate(document.summary, start=_SUMMARY_FIRST_DATA_ROW):
        worksheet.cell(row=offset, column=1, value=label)
        worksheet.cell(row=offset, column=2, value=value)
    worksheet.column_dimensions["A"].width = 28
    worksheet.column_dimensions["B"].width = 40


def _append_inline_tables(
    worksheet: "Worksheet",
    document: ReportDocument,
    *,
    bold: "Font",
    title_font: "Font",
    start_row: int,
) -> None:
    """Mirror the continuous PDF layout on the summary worksheet."""
    from openpyxl.styles import Alignment

    next_row = start_row
    for table in document.tables:
        worksheet.cell(row=next_row, column=1, value=table.name).font = title_font
        next_row += 1
        for col_idx, column in enumerate(table.columns, start=1):
            cell = worksheet.cell(row=next_row, column=col_idx, value=column)
            cell.font = bold
            cell.alignment = Alignment(wrap_text=True)
        next_row += 1
        for row in table.rows:
            for col_idx, cell_value in enumerate(row, start=1):
                worksheet.cell(row=next_row, column=col_idx, value=excel_cell(cell_value))
            next_row += 1
        next_row += 2


def _write_detail_sheet(worksheet: "Worksheet", table: ReportTable, bold: "Font") -> None:
    from openpyxl.styles import Alignment

    worksheet.freeze_panes = "A2"
    last_column = excel_column_letter(len(table.columns))
    worksheet.auto_filter.ref = f"A1:{last_column}{max(len(table.rows) + 1, 1)}"

    for col_idx, column in enumerate(table.columns, start=1):
        cell = worksheet.cell(row=1, column=col_idx, value=column)
        cell.font = bold
        cell.alignment = Alignment(wrap_text=True)
    for row_idx, row in enumerate(table.rows, start=2):
        for col_idx, cell_value in enumerate(row, start=1):
            worksheet.cell(row=row_idx, column=col_idx, value=excel_cell(cell_value))
    for col_idx in range(1, len(table.columns) + 1):
        letter = worksheet.cell(row=1, column=col_idx).column_letter
        worksheet.column_dimensions[letter].width = _DETAIL_COLUMN_WIDTH


def render_excel(document: ReportDocument) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    summary_ws = workbook.active
    summary_ws.title = "Summary"

    bold = Font(bold=True)
    title_font = Font(bold=True, size=14)

    _write_header_block(summary_ws, document, title_font)
    _write_summary_block(summary_ws, document, bold)
    _append_inline_tables(
        summary_ws,
        document,
        bold=bold,
        title_font=title_font,
        start_row=_SUMMARY_FIRST_DATA_ROW + len(document.summary) + 2,
    )
    summary_ws.freeze_panes = "A8"

    used_names: set[str] = {"Summary"}
    for table in document.tables:
        sheet = workbook.create_sheet(unique_sheet_name(table.name, used_names))
        _write_detail_sheet(sheet, table, bold)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


__all__ = ["render_excel"]
